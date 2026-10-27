"""Team installations without the Codex artifact runtime use this portable XLSX writer."""
import io
import math


def build_workbook(payload):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment,Border,Side,Font,PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ValueError('Excel 依赖未安装，请运行 Skill 的 bootstrap.py --install。') from exc
    mode=payload.get('template','archive');wb=Workbook();ws=wb.active
    ws.title={'archive':'详细归档','script':'脚本文案表','shooting':'拍摄需求表'}[mode]
    widths={'script':[24,115,36,20,28],'shooting':[22,19,32,70,32,32,24,26],'archive':[8,32,100,10,20,32,23,16]}[mode]
    ws.sheet_view.showGridLines=False;ws.freeze_panes='A2'
    rows=[payload['headers'],*payload['rows']];border=Border(**{edge:Side(style='thin',color='D0D0D0') for edge in ('left','right','top','bottom')})
    body_column={'script':2,'shooting':4,'archive':3}[mode]
    for r,row in enumerate(rows,1):
        for col,value in enumerate(row,1):
            cell=ws.cell(r,col,value)
            # Generated copy is text, even when its first character is '='.
            if isinstance(value,str):cell.data_type='s'
            cell.font=Font(name='Arial',size=11 if r==1 else 10,bold=r==1,color='1D1D1F')
            cell.fill=PatternFill('solid',fgColor='FFFFFF');cell.border=border
            cell.alignment=Alignment(horizontal='left' if r>1 and col==body_column else 'center',vertical='center',wrap_text=True)
        if r==1:ws.row_dimensions[r].height=42
        else:
            text=str(row[body_column-1] or '')
            lines=sum(max(1,math.ceil(len(line)/(35 if mode=='shooting' else 55))) for line in text.split('\n'))
            ws.row_dimensions[r].height=min(409,max(160 if mode=='shooting' else 130,lines*13+24))
    for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
    ws.auto_filter.ref=ws.dimensions
    output=io.BytesIO();wb.save(output);return output.getvalue()
