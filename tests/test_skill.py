import copy
from http.server import BaseHTTPRequestHandler,HTTPServer
import io
import json
from pathlib import Path
import sys
import tempfile
import threading
import unittest
import zipfile

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'));sys.path.insert(0,str(ROOT/'assets/studio'))
from studio_agent import Studio,selected_batch,save_output,invoke
from install_skill import install
from client.xlsx_portable import build_workbook
from client.draft_exports import table_payload,DraftTextExport
from prepare_models import model_status,REQUIRED,MODEL_CONTRACT
from bootstrap import write_launcher


class LocalTransport(unittest.TestCase):
    def setUp(self):
        test=self
        class Handler(BaseHTTPRequestHandler):
            def log_message(self,*args):pass
            def do_GET(self):
                self.send_response(200)
                if self.path=='/':self.send_header('Set-Cookie','studio_session=test; HttpOnly; SameSite=Strict')
                self.end_headers();self.wfile.write(b'{}' if self.path=='/' else b'{"ok":true}')
            def do_POST(self):
                test.assertEqual(self.headers['Cookie'],'studio_session=test');test.assertEqual(self.headers['X-Studio-Request'],'1')
                body=self.rfile.read(int(self.headers['Content-Length']))
                test.assertIn(b'local-voice-bytes',body);test.assertIn(b'name="auto_ai"',body)
                self.send_response(202);self.end_headers();self.wfile.write(b'{"id":"test-job"}')
        self.server=HTTPServer(('127.0.0.1',0),Handler)
        self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start()
        self.addCleanup(self.server.server_close);self.addCleanup(self.server.shutdown)
        self.url='http://127.0.0.1:'+str(self.server.server_address[1])

    def test_cookie_and_streamed_upload(self):
        with tempfile.TemporaryDirectory() as tmp:
            media=Path(tmp)/'示例.mp3';media.write_bytes(b'local-voice-bytes'*50)
            result=Studio(self.url).upload(media,'已确认资料',False)
            self.assertEqual(result,{'id':'test-job'})

    def test_external_service_is_rejected_before_connecting(self):
        with self.assertRaises(ValueError):Studio('https://example.com/')


class PackagingContract(unittest.TestCase):
    def test_script_and_document_cli_routes_and_exclusive_sources(self):
        class Client:
            def request(self,path,*args):return {'path':path,'args':args}
            def upload(self,*args,**kwargs):return {'upload':args,'options':kwargs}
            def upload_materials(self,*args,**kwargs):return {'materials':args,'options':kwargs}
        client=Client()
        self.assertEqual(invoke(client,{'action':'import-script','text':'脚本原文'})['path'],'/jobs/script')
        self.assertTrue(invoke(client,{'action':'import-script','file':'script.docx'})['options']['script'])
        result=invoke(client,{'action':'extract-material','job_id':'j','document':'product.pdf'})
        self.assertEqual(result['materials'],('j',['product.pdf'],'switch'));self.assertTrue(result['options']['document'])
        with self.assertRaises(ValueError):invoke(client,{'action':'import-script','file':'script.docx','text':'不能混用'})
        with self.assertRaises(ValueError):invoke(client,{'action':'extract-material','job_id':'j','text':'不能混用','document':'product.pdf'})

    def test_export_prefers_current_file_backed_manuscript(self):
        job={'batches':[{'id':'b','variants':[{'segments':[{'text':'生成原稿'}],
            'manuscript':{'revision':2,'text':'已修改成稿'}}]}]}
        batch,items=selected_batch(job,{})
        self.assertEqual(items,[{'index':0,'text':'已修改成稿','revision':2}])
        _,explicit=selected_batch(job,{'edits':{'1':'一次性导出稿'}})
        self.assertEqual(explicit,[{'index':0,'text':'一次性导出稿'}])

    def test_cli_uses_one_based_script_index_and_requires_revision(self):
        class Client:
            def request(self,path,*args):
                if path=='/jobs/j':return {'batches':[{'id':'b','variants':[{}]}]}
                return {'path':path,'args':args}
        self.assertEqual(invoke(Client(),{'action':'manuscript','job_id':'j'})['path'],'/jobs/j/drafts/b/0/versions')
        with self.assertRaises(ValueError):invoke(Client(),{'action':'rewrite','job_id':'j'})
        with self.assertRaises(ValueError):invoke(Client(),{'action':'manuscript','job_id':'j','script_index':0})

    def test_missing_model_invalidates_success_receipt(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);weight=root/'model.pt';weight.write_bytes(b'verified-checkpoint')
            receipt=root/'models-ready.json'
            data={'contract':MODEL_CONTRACT,'cache':str(root),'loaded':dict.fromkeys(REQUIRED,True),
                  'files':[{'path':str(weight),'bytes':weight.stat().st_size}]}
            receipt.write_text(json.dumps(data))
            self.assertTrue(model_status(receipt,root))
            weight.write_bytes(b'truncated');self.assertFalse(model_status(receipt,root))
            weight.unlink();self.assertFalse(model_status(receipt,root))

    def test_launcher_supports_unicode_and_quotes_in_runtime_path(self):
        import ast
        with tempfile.TemporaryDirectory(prefix="文案'空 格") as tmp:
            root=Path(tmp);launcher,_=write_launcher(root,root/'python.exe')
            module=ast.parse(launcher.read_text(encoding='utf-8'))
            assigned=next(n for n in module.body if isinstance(n,ast.Assign))
            command=ast.literal_eval(assigned.value)
            self.assertIn(str(root),command)
            self.assertEqual(command[-2:],['start','--open'])
            self.assertTrue((root/'studio.cmd').exists())

    def test_text_false_never_authorizes_cloud_write_or_overwrite(self):
        for flag in ('auto_ai','write','overwrite','quality_check'):
            with self.assertRaises(ValueError):invoke(None,{'action':'feishu',flag:'false'})

    def test_export_selection_never_mixes_batches_and_preserves_edits(self):
        old={'id':'old','variants':[{'segments':[{'text':'旧稿'}]}]}
        latest={'id':'new','variants':[{'segments':[{'text':'第一段'},{'text':'第二段'}]},{'segments':[{'text':'不选'}]}]}
        batch,items=selected_batch({'batches':[old,latest]},{'indices':[1],'edits':{'1':'修改正文'}})
        self.assertEqual(batch['id'],'new');self.assertEqual(items,[{'index':0,'text':'修改正文'}])
        with self.assertRaises(ValueError):selected_batch({'batches':[latest]},{'indices':[0]})

    def test_output_requires_explicit_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'output.txt';save_output(path,b'old')
            with self.assertRaises(ValueError):save_output(path,b'new')
            self.assertEqual(path.read_bytes(),b'old')

    def test_install_upgrade_keeps_backup(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp=Path(tmp);source=tmp/'source';source.mkdir();(source/'SKILL.md').write_text('old')
            first=install(source,tmp/'agent-skills');target=Path(first['installed_path']);self.assertEqual((target/'SKILL.md').read_text(),'old')
            with self.assertRaises(ValueError):install(source,tmp/'agent-skills')
            (source/'SKILL.md').write_text('new')
            updated=install(source,tmp/'agent-skills',True)
            self.assertEqual((target/'SKILL.md').read_text(),'new')
            self.assertEqual((Path(updated['backup_path'])/'SKILL.md').read_text(),'old')
            self.assertFalse(Path(updated['backup_path']).is_relative_to(tmp/'agent-skills'))
            self.assertEqual(len(list((tmp/'agent-skills').glob('*/SKILL.md'))),1)

    def test_portable_excel_has_full_text_no_formula_and_expected_columns(self):
        from openpyxl import load_workbook
        item={'id':'demo','title':'参考','kind':'upload','batches':[{'id':'batch','mode':'same','created':0,'variants':[{'title':'稿件','segments':[{'text':'原稿'}]}]}]}
        for template,count in [('script',5),('shooting',8),('archive',8)]:
            request=DraftTextExport(batch_id='batch',format='xlsx',template=template,items=[{'index':0,'text':'=不是公式\n\n保留段落'}])
            before=copy.deepcopy(item);raw=build_workbook(table_payload(item,request));self.assertEqual(item,before)
            wb=load_workbook(io.BytesIO(raw));ws=wb.active
            self.assertEqual(ws.max_column,count);self.assertEqual(ws.freeze_panes,'A2');self.assertTrue(ws.auto_filter.ref)
            if template!='shooting':
                cell=ws.cell(2,2 if template=='script' else 3)
                self.assertEqual(cell.value,'=不是公式\n\n保留段落');self.assertEqual(cell.data_type,'s')
            else:self.assertIsNone(ws.cell(2,2).value)


if __name__=='__main__':unittest.main()
